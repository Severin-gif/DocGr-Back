import importlib.util
import io
import pathlib
import tempfile
import unittest
import zipfile
from pypdf import PdfReader
from PIL import Image
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import load_workbook

spec=importlib.util.spec_from_file_location('engine',pathlib.Path(__file__).with_name('document-tools.py'))
engine=importlib.util.module_from_spec(spec);spec.loader.exec_module(engine)

class ToolsTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=pathlib.Path(self.tmp.name)
        self.pdf=self.root/'source.pdf';c=canvas.Canvas(str(self.pdf),pagesize=(595,842))
        for text in ['Договор SECRET 123','Приложение 456']:
            c.setFont('DG',16);c.drawString(50,760,text);c.showPage()
        c.save()
    def tearDown(self):self.tmp.cleanup()
    def call(self,tool,options=None,files=None):
        work=self.root/('work-'+tool);work.mkdir(exist_ok=True)
        result=engine.process({'tool':tool,'options':options or {},'files':[{'path':str(p),'title':p.name} for p in (files or [self.pdf])]},work)
        return result,work/result['outputs'][0]['file'] if 'outputs' in result else None
    def test_merge(self):
        _,p=self.call('merge',files=[self.pdf,self.pdf]);self.assertEqual(len(PdfReader(p).pages),4)
    def test_reorder_rotate(self):
        _,p=self.call('pages',{'pages':'2,1','rotation':90});r=PdfReader(p);self.assertIn('456',r.pages[0].extract_text());self.assertEqual(r.pages[0].rotation,90)
    def test_delete_split_and_invalid(self):
        _,p=self.call('pages',{'pages':'1','mode':'delete'});self.assertEqual(len(PdfReader(p).pages),1)
        _,p=self.call('pages',{'mode':'split'});self.assertEqual(len(zipfile.ZipFile(p).namelist()),2)
        with self.assertRaises(ValueError):self.call('pages',{'pages':'9'})
        with self.assertRaises(ValueError):self.call('pages',{'mode':'delete','pages':'1-2'})
    def test_number_watermark(self):
        _,p=self.call('number',{'pages':'2','start':5});r=PdfReader(p);self.assertTrue(r.pages[1].extract_text().strip().endswith('5'))
        _,p=self.call('watermark',{'text':'Черновик'});self.assertIn('Черновик',PdfReader(p).pages[0].extract_text())
    def test_protect(self):
        _,p=self.call('protect',{'password':'Secret-123'});r=PdfReader(p);self.assertTrue(r.is_encrypted);self.assertFalse(r.decrypt('wrong'));self.assertTrue(r.decrypt('Secret-123'));self.assertEqual(len(r.pages),2)
    def test_compress(self):
        _,p=self.call('compress');self.assertLessEqual(p.stat().st_size,self.pdf.stat().st_size);self.assertIn('SECRET',PdfReader(p).pages[0].extract_text())
    def test_images(self):
        img=self.root/'photo.png';Image.new('RGB',(200,300),'green').save(img)
        _,p=self.call('images-pdf',{'landscape':True},[img,img]);r=PdfReader(p);self.assertEqual(len(r.pages),2);self.assertGreater(r.pages[0].mediabox.width,r.pages[0].mediabox.height)
    def test_ocr_existing_text(self):
        _,p=self.call('ocr');self.assertIn('Договор',PdfReader(p).pages[0].extract_text())
    def test_scan_ocr(self):
        img=engine.image_page(self.pdf,0,self.root,180);scan=self.root/'scan.pdf';engine.image_pdf([img],scan)
        _,p=self.call('ocr',{'languages':'eng'},[scan]);self.assertIn('SECRET',PdfReader(p).pages[0].extract_text())
    def test_docx_conversion(self):
        doc=Document();doc.add_paragraph('Document contract 123');source=self.root/'contract.docx';doc.save(source)
        _,p=self.call('docx-pdf',files=[source]);self.assertIn('contract',PdfReader(p).pages[0].extract_text())
        _,p=self.call('pdf-docx');self.assertIn('Договор', '\n'.join(x.text for x in Document(p).paragraphs))
    def test_tables_formula_safety(self):
        table=self.root/'table.pdf';c=canvas.Canvas(str(table),pagesize=(595,842));c.setFont('DG',12)
        for x in [40,240,440]:c.line(x,600,x,750)
        for y in [600,650,700,750]:c.line(40,y,440,y)
        for y,a,b in [(720,'Name','Value'),(670,'=1+1','100'),(620,'00123','200')]:c.drawString(50,y,a);c.drawString(250,y,b)
        c.save();_,p=self.call('pdf-xlsx',files=[table]);wb=load_workbook(p);self.assertEqual(wb.active['A2'].value,'=1+1');self.assertEqual(wb.active['A2'].data_type,'s');self.assertEqual(wb.active['A3'].value,'00123')
    def test_redaction_removes_text(self):
        result,p=self.call('redact',{'terms':'SECRET'});r=PdfReader(p);self.assertNotIn('SECRET',''.join(x.extract_text() for x in r.pages));self.assertTrue(result['warnings']);self.assertNotIn(b'SECRET',p.read_bytes())
        with self.assertRaises(ValueError):self.call('redact',{'terms':'NOT PRESENT'})
    def test_redaction_rect(self):
        _,p=self.call('redact',{'rects':[{'page':1,'x':.05,'y':.03,'w':.8,'h':.15}]});im=engine.image_page(p,0,self.root,72);self.assertEqual(im.getpixel((80,60)),(0,0,0))
    def test_compare(self):
        result,p=self.call('compare',files=[self.pdf,self.pdf]);self.assertEqual(len(PdfReader(p).pages),2);self.assertIn('не обнаружены',result['warnings'][0])
    def test_bundle_and_zip(self):
        _,p=self.call('bundle',files=[self.pdf,self.pdf]);r=PdfReader(p);self.assertEqual(len(r.pages),5);self.assertIn('Опись',r.pages[0].extract_text())
        _,p=self.call('bundle',{'format':'zip'},[self.pdf,self.pdf]);names=zipfile.ZipFile(p).namelist();self.assertEqual(len(names),3);self.assertIn('Опись.txt',names)
    def test_inspect(self):
        result,_=self.call('inspect',{'page':2});self.assertEqual(result['pages'],2);self.assertTrue(result['image'].startswith('data:image/jpeg;'))

if __name__=='__main__':unittest.main()
